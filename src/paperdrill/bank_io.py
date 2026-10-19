# -*- coding: utf-8 -*-
"""题库读写：支持 xlsx/xlsm、csv、txt、json、docx、pdf。

设计要点
--------
* 表格型数据源（Excel/CSV/TXT/JSON、Word 表格、PDF 表格）走 ``rows_to_questions``：
  列头同义词识别 + 位置回退，兼容用户自行整理过的表格；
* 段落型数据源（Word 段落、PDF 文本）走 ``lines_to_questions``：
  识别章节标题（题型）、题干、选项、答案行、解析行，答案可嵌在题干括号里；
* 逐行校验，坏行不进入题库，但完整记录原因供界面提示；
* 全部读入为 ``list[Question]``，交给领域层处理。
"""
from __future__ import annotations

import csv
import io
import json
import re
from collections import Counter, deque
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Sequence

from .models import (
    OPTION_LABELS,
    AnswerParseError,
    QType,
    Question,
    normalize_display,
    normalize_text,
    parse_answer,
    parse_qtype,
)

__all__ = [
    "BankIssue", "BankLoadResult", "TextLine", "load_bank", "questions_to_rows",
    "export_questions_to_xlsx", "write_template_xlsx", "build_sample_questions",
    "bank_stats", "SUPPORTED_SUFFIXES", "lines_to_questions",
    "TEXT_SUFFIXES", "WORD_SUFFIXES", "PDF_SUFFIXES",
]

TEXT_SUFFIXES = (".csv", ".txt")
WORD_SUFFIXES = (".docx", ".docm")
PDF_SUFFIXES = (".pdf",)
SUPPORTED_SUFFIXES = (".xlsx", ".xlsm") + TEXT_SUFFIXES + (".json",) + WORD_SUFFIXES + PDF_SUFFIXES

# ---------------------------------------------------------------- 列头同义词
_HEADER_SYNONYMS: dict[str, list[str]] = {
    "type": ["题型", "题目类型", "类型", "试题类型", "题目种类",
             "type", "kind", "category"],
    "stem": ["题目标题", "题目", "题干", "试题", "问题", "题目内容", "试题内容", "试题题目",
             "题干内容", "题干题目", "题目描述", "题目正文", "试题正文", "试题题干", "内容",
             "stem", "question", "title", "questiontext", "description"],
    "analysis": ["解析", "答案解析", "试题解析", "题目解析", "说明", "分析", "参考答案解析",
                 "备注", "依据", "来源", "知识点", "analysis", "explain", "explanation"],
    "answer": ["答案", "正确答案", "参考答案", "标准答案", "正确选项",
               "试题答案", "题目答案",
               "answer", "key", "correctanswer"],
}

# 「分类 / 题目种类 / 试题分类」只作为**没有**「题型/类型」列时的兜底同义词，
# 避免「分类」排在「题型」前面时抢走题型列（审计 R5）。
_TYPE_FALLBACK_TOKENS = ("分类", "题目种类", "试题分类")
_OPTION_HEADER_RE = re.compile(r"^(?:选项)?\s*([A-Ha-h])\s*(?:选项)?$")
_HEADER_SCAN_ROWS = 20                                   # 列头行扫描范围（原为 5 行）


@dataclass
class BankIssue:
    """一条无法导入的记录。"""

    row: int
    reason: str
    raw: str = ""

    def __str__(self) -> str:  # pragma: no cover - 展示用
        return f"第 {self.row} 行：{self.reason}"


@dataclass
class BankLoadResult:
    """题库加载结果。"""

    questions: list[Question] = field(default_factory=list)
    issues: list[BankIssue] = field(default_factory=list)
    source: str = ""
    sheet: str = ""
    header: list[str] = field(default_factory=list)
    column_map: dict[str, object] = field(default_factory=dict)

    @property
    def ok_count(self) -> int:
        return len(self.questions)

    def type_counts(self) -> dict[QType, int]:
        counts = {QType.SINGLE: 0, QType.MULTIPLE: 0, QType.JUDGE: 0}
        for q in self.questions:
            counts[q.qtype] += 1
        return counts


# ---------------------------------------------------------------- 行数据读取
def _read_xlsx(path: Path) -> tuple[str, list[list[str]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    raw_wb = None
    try:
        raw_wb = load_workbook(path, data_only=False, read_only=True)
    except Exception:
        raw_wb = None
    try:
        candidates: list[tuple[int, int, int, str, list[list[str]]]] = []
        for order, ws in enumerate(wb.worksheets):
            rows: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                rows.append([normalize_display(c) for c in row])
            if not rows:
                continue
            header_idx, _, _ = _detect_header(rows)
            # 优先级：能识别列头 > 行数多 > 工作表靠前
            candidates.append((1 if header_idx >= 0 else 0, len(rows), -order, ws.title, rows))
        if not candidates:
            return "", []
        best = max(candidates, key=lambda t: (t[0], t[1], t[2]))
        title, rows = best[3], best[4]
        # 公式单元格若没有缓存值，data_only=True 读出来是 None；标记出来，
        # 让"答案为空"变成可操作的提示（审计 R24）。
        if raw_wb is not None and title in raw_wb.sheetnames:
            raw_rows = list(raw_wb[title].iter_rows(values_only=True))
            for r, row in enumerate(rows):
                if r >= len(raw_rows):
                    break
                raw_row = raw_rows[r]
                for c, value in enumerate(row):
                    if value or c >= len(raw_row):
                        continue
                    raw_value = raw_row[c]
                    if isinstance(raw_value, str) and raw_value.startswith("="):
                        row[c] = f"[公式无缓存值] {raw_value}"
        return title, rows
    finally:
        wb.close()
        if raw_wb is not None:
            raw_wb.close()


def _decode(data: bytes) -> str:
    # Excel「Unicode 文本」/ 记事本「UTF-16 LE」导出带 BOM，必须先按 UTF-16 解
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError:
            pass
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")  # pragma: no cover


_DELIM_CANDIDATES = ("\t", "|", ";", ",", "，")


def _sniff_delimiter(text: str) -> str:
    """猜测列分隔符。

    题库单元格里经常出现 ``;``（单列多选项，如 ``A.甲;B.乙``），所以不能
    「谁先出现就用谁」——那样会把 ``;`` 当成列分隔符，整张表错列。
    这里改为：候选分隔符中取**各行列数最一致**的一个，一致程度相同时取列数更多者。
    """
    lines = [ln for ln in text.splitlines()[:50] if ln.strip()]
    best_delim, best_score = "", (0, 0)
    for delim in _DELIM_CANDIDATES:
        counts = [ln.count(delim) for ln in lines]
        if not any(counts):
            continue
        modal_count, agree = Counter(counts).most_common(1)[0]
        if modal_count == 0:            # 该分隔符只出现在少数行 → 多半是单元格内容
            continue
        score = (agree, modal_count)
        if score > best_score:
            best_delim, best_score = delim, score
    if best_delim:
        return best_delim
    try:
        return csv.Sniffer().sniff(text[:4096]).delimiter
    except Exception:
        return ","


def _read_delimited(path: Path) -> list[list[str]]:
    text = _decode(path.read_bytes())
    if not text.strip():
        return []
    delim = _sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    try:
        return [[normalize_display(c) for c in row] for row in reader]
    except csv.Error as exc:
        raise ValueError(
            f"文本题库解析失败：{exc}。请确认文件是 UTF-8/GBK 的 CSV/TXT"
            "（Excel「Unicode 文本」可直接导入；若仍失败请另存为 UTF-8）") from exc


def _read_json(path: Path) -> list[list[str]]:
    """JSON 支持 {"questions":[...]} 或直接数组；元素为对象。"""
    data = json.loads(_decode(path.read_bytes()) or "[]")
    if isinstance(data, dict):
        for key in ("questions", "data", "items", "list", "题目", "题库"):
            if key in data and isinstance(data[key], list):
                data = data[key]
                break
        else:
            raise ValueError("JSON 结构无法识别：未找到题目数组（questions/data/items）")
    if not isinstance(data, list):
        raise ValueError("JSON 顶层必须是数组或包含 questions 数组的对象")

    def pick(obj: dict, names: Sequence[str], default=""):
        for n in names:
            if n in obj and obj[n] not in (None, ""):
                return obj[n]
        lowered = {str(k).lower(): v for k, v in obj.items()}
        for n in names:
            if n.lower() in lowered and lowered[n.lower()] not in (None, ""):
                return lowered[n.lower()]
        return default

    rows: list[list[str]] = []
    for obj in data:
        if not isinstance(obj, dict):
            rows.append([normalize_display(obj)])
            continue
        opts = pick(obj, ["options", "选项", "选项列表", "choices"], [])
        if isinstance(opts, str):
            parts = [p for p in re.split(r"[\n\r]+", opts) if p.strip()]
            if len(parts) == 1:
                # 单行 "A.甲 B.乙 C.丙" → 按选项标记拆开（审计 S7）
                parsed = _split_options(parts[0])[0]
                if len(parsed) >= 2:
                    opts = [text for _letter, text in parsed]
                else:
                    opts = parts
            else:
                opts = parts
        if isinstance(opts, dict):
            opts = [opts[k] for k in sorted(opts)]
        opt_texts = [normalize_display(o) for o in (opts or [])]
        # 选项文本自带 "A." 前缀时去掉，避免与界面标签重复（A A. 甲）
        for idx, text in enumerate(opt_texts):
            m = re.match(r"^\s*([A-Ha-hＡ-Ｈａ-ｈ])\s*[.．、)）]\s*(.+)$", text)
            if m and _letter(m.group(1)) == OPTION_LABELS[idx]:
                opt_texts[idx] = m.group(2).strip()
        # 允许 A/B/C/D 与 选项A 两种键
        for letter in OPTION_LABELS:
            v = pick(obj, [f"选项{letter}", f"option{letter}", letter, letter.lower()], "")
            if v not in ("", None):
                while len(opt_texts) <= OPTION_LABELS.index(letter):
                    opt_texts.append("")
                opt_texts[OPTION_LABELS.index(letter)] = normalize_display(v)
        while opt_texts and not opt_texts[-1]:
            opt_texts.pop()
        opt_texts = (opt_texts + [""] * 8)[:8]
        rows.append([
            normalize_display(pick(obj, ["type", "题型", "题目类型", "类型", "试题类型"], "")),
            normalize_display(pick(obj, ["stem", "question", "题干", "题目", "title",
                                         "题目标题", "题目内容", "试题内容", "试题题目"], "")),
            *opt_texts,
            normalize_display(pick(obj, ["analysis", "解析", "说明", "答案解析", "依据", "来源"], "")),
            normalize_display(pick(obj, ["answer", "答案", "正确答案", "参考答案", "标准答案", "key"], "")),
        ])
    return rows


# ---------------------------------------------------------------- 列映射
def _header_token(cell: object) -> str:
    """列头比对用的归一化文本：全半角统一、去空白与换行。"""
    return normalize_text(cell).lower().replace(" ", "").replace("\n", "").replace("\t", "")


def _option_letter_of_header(cell: object) -> "str | None":
    """把列头单元格解析为选项字母 A–H；不是选项列返回 None。

    兼容 `选项A` / `A选项` / `选项 A` / `A` / `A、` / `选项（A）` / `选项1`
    / `optionA` / `opt A` 等写法。
    """
    token = _header_token(cell)
    if not token:
        return None
    for affix in ("选项", "options", "option", "opt"):
        token = token.replace(affix, "")
    token = token.strip("()[]（）第项列：:、,，.")
    if len(token) != 1:
        return None
    if token in "abcdefgh":
        return token.upper()
    if token in "12345678":
        return "ABCDEFGH"[int(token) - 1]
    return None


def _detect_header(rows: list[list[str]]) -> tuple[int, dict[str, object], list[str]]:
    """在前若干行中寻找列头行，返回 (行索引, 列映射, 原始列头)。

    识别条件：该行含题干同义词，且**至少有 2 个选项字母列**，或**有 1 个"选项"合并列**
    （内容形如 `A.甲 B.乙`，后续会自动展开）。
    """
    for i, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        mapping: dict[str, object] = {}
        opt_cols: dict[int, str] = {}
        combo_col: "int | None" = None
        fallback_type: "int | None" = None
        for j, cell in enumerate(row):
            token = _header_token(cell)
            if not token:
                continue
            if "type" not in mapping and token in _TYPE_FALLBACK_TOKENS:
                # 只记录兜底列位置，绝不抢走真正的「题型/类型」列
                if fallback_type is None:
                    fallback_type = j
                continue
            matched = False
            for key, names in _HEADER_SYNONYMS.items():
                if token in [_header_token(n) for n in names]:
                    mapping.setdefault(key, j)
                    matched = True
                    break
            if matched:
                continue
            if token in _COMBINED_OPTION_HEADERS:
                if combo_col is None:
                    combo_col = j
                continue
            letter = _option_letter_of_header(cell)
            if letter:
                opt_cols[j] = letter
        if "stem" not in mapping:
            continue
        if "type" not in mapping and fallback_type is not None:
            mapping["type"] = fallback_type
        # 有 2 个以上选项列、或有 1 个「选项」合并列、或至少有题型/答案列 → 认为是列头行
        if (len(opt_cols) >= 2 or (combo_col is not None and not opt_cols)
                or "type" in mapping or "answer" in mapping):
            mapping["options"] = opt_cols
            if combo_col is not None:
                mapping["combined_option_col"] = combo_col
            mapping["header_row"] = i
            return i, mapping, list(row)
    return -1, {}, []


def _positional_map(width: int) -> "dict[str, object] | None":
    """无列头时的位置回退策略；**仅为与常见模板吻合的列数**回退。

    12 列（标准全宽）、8 列（题型/题干/A–D/解析/答案）、3 列（题型/题干/答案）以外
    一律返回 ``None``：宁可明确报错，也不要按位置猜错列（例如把 D 选项当成答案）。
    """
    if width >= 12:
        return {"type": 0, "stem": 1, "options": {j: OPTION_LABELS[j - 2] for j in range(2, 10)},
                "analysis": 10, "answer": 11}
    if width == 8:
        return {"type": 0, "stem": 1, "options": {j: OPTION_LABELS[j - 2] for j in range(2, 6)},
                "analysis": 6, "answer": 7}
    if width == 3:
        return {"type": 0, "stem": 1, "options": {}, "analysis": -1, "answer": 2}
    return None


# 表格单元格里的手动折行：汉字（含中文标点）之间直接相接，其它情况补一个空格。
# 段落型题库的换行才有结构含义，表格里出现的换行只是整理时的折行。
_CJK_WRAP_RE = re.compile(
    r"(?<=[\u3000-\u303f\u4e00-\u9fff\u3400-\u4dbf\uf900-\ufaff])[ \t]*\n[ \t]*"
    r"(?=[\u3000-\u303f\u4e00-\u9fff\u3400-\u4dbf\uf900-\ufaff])"
)


def _unwrap_cell(text: str) -> str:
    """把单元格内的手动折行拼回一行：汉字之间直接相接，其它情况补一个空格。"""
    if "\n" not in text:
        return text
    return re.sub(r"[ \t]*\n[ \t]*", " ", _CJK_WRAP_RE.sub("", text))


def _cell(row: Sequence[str], idx: object) -> str:
    """取表格单元格文本（顺带把手动折行拼回一行）。"""
    if idx is None:
        return ""
    try:
        i = int(idx)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return ""
    if i < 0 or i >= len(row):
        return ""
    return normalize_display(_unwrap_cell(str(row[i])))


# ---------------------------------------------------------------- 主流程
def _infer_qtype(answer: str, options: list[str]) -> "QType | None":
    """缺「题型」列时由答案形态推断题型（与段落解析器同一套规则）。"""
    token = normalize_text(answer)
    if not token:
        return None
    if _judge_from_token(token) is not None and (not options or len(options) < 2):
        return QType.JUDGE
    letters = _letters_from_token(token)
    if len(letters) >= 2:
        return QType.MULTIPLE
    if len(letters) == 1:
        return QType.JUDGE if not options else QType.SINGLE
    if _judge_from_token(token) is not None:
        return QType.JUDGE
    return None


def _option_subheader_letters(row: Sequence[str]) -> "dict[int, str] | None":
    """两行表头：若该行非空单元格**全部**是选项字母/「选项A」子表头，返回列映射。

    只认 >=2 个选项列且没有其它表头文本，避免把普通数据行误当子表头。
    """
    letters: dict[int, str] = {}
    others = 0
    for j, cell in enumerate(row):
        token = _header_token(cell)
        if not token:
            continue
        letter = _option_letter_of_header(cell)
        if letter:
            letters[j] = letter
        else:
            others += 1
    if len(letters) >= 2 and others == 0:
        return letters
    return None


def rows_to_questions(rows: list[list[str]], *, start_row: int = 1
                      ) -> tuple[list[Question], list[BankIssue], list[str], dict]:
    """把二维表行转换为 Question 列表。"""
    issues: list[BankIssue] = []
    questions: list[Question] = []
    if not rows:
        return questions, issues, [], {}

    header_idx, cmap, header = _detect_header(rows)
    # 两行表头（第一行「选项」组标题 / 第二行 A B C D 子表头）：把子表头并入列映射
    if header_idx >= 0 and header_idx + 1 < len(rows):
        sub_letters = _option_subheader_letters(rows[header_idx + 1])
        if sub_letters is not None:
            cmap["options"] = sub_letters
            cmap.pop("combined_option_col", None)
            header_idx += 1
            header = list(rows[header_idx])
    if header_idx >= 0 and cmap.get("combined_option_col") is not None:
        # 「选项」在一列里（如 `A.甲 B.乙 C.丙 D.丁`）→ 展开成 选项A…选项H 再重新识别
        expanded = _expand_combined_options(rows)
        if expanded is not rows:
            header_idx, cmap, header = _detect_header(expanded)
            rows = expanded

    if header_idx >= 0:
        data_rows = rows[header_idx + 1:]
        first_data_row = start_row + header_idx + 1
        width = max((len(r) for r in rows), default=0)
    else:
        data_rows = rows
        first_data_row = start_row
        width = max((len(r) for r in rows), default=0)
        cmap = _positional_map(width)
        if width == 3:
            sample = next((r for r in rows if any(normalize_text(c) for c in r)), None)
            if sample is not None:
                first_cell = _cell(sample, 0)
                if parse_qtype(first_cell) is None and _cell(sample, 2):
                    # 无表头 3 列更常见的是「题干 / 选项 / 答案」，而不是「题型/题干/答案」
                    cmap = {"stem": 0, "answer": 2, "combined_option_col": 1}
        if cmap is None:
            raise ValueError(
                f"未找到可识别的列头行（需包含「题干」以及至少 2 个「选项A…」列，或 1 个「选项」合并列）；"
                f"当前表格 {width} 列且无法按位置安全推断。请在第一行补上列头，例如："
                "题型 | 题目标题 | 选项A | 选项B | 选项C | 选项D | 解析 | 答案")
        header = []

    opt_cols: dict[int, str] = dict(cmap.get("options", {}))  # type: ignore[arg-type]

    running = 0
    for offset, row in enumerate(data_rows):
        row_no = first_data_row + offset
        if not any(normalize_text(c) for c in row):
            continue
        raw_preview = " | ".join(normalize_text(c) for c in row[:4])[:120]

        texts: list[tuple[str, str]] = []
        combo_col = cmap.get("combined_option_col")
        if combo_col is not None and not opt_cols:
            texts = sorted(_split_option_cell(_cell(row, combo_col)))
            texts = [(letter, text) for letter, text in texts if text]
        else:
            for col, letter in sorted(opt_cols.items(), key=lambda kv: kv[1]):
                text = _cell(row, col)
                texts.append((letter, text))
            texts.sort(key=lambda kv: kv[0])
            while texts and not texts[-1][1]:
                texts.pop()
        options_preview = [t for _, t in texts]

        raw_type = _cell(row, cmap.get("type"))
        raw_answer = _cell(row, cmap.get("answer"))
        if raw_answer.startswith("[公式无缓存值]"):
            issues.append(BankIssue(
                row_no,
                "答案单元格是公式且没有缓存值：请在 Excel 中复制答案列 → 选择性粘贴为「值」，"
                "或把公式结果转成纯文本后重新导入", raw_preview))
            continue
        qtype = parse_qtype(raw_type)
        if qtype is None:
            if raw_type.strip():
                # 「选择题 / 客观题」等泛化题型仍按"无法识别"处理（与既有测试/产品约定一致）
                issues.append(BankIssue(row_no, f"题型无法识别：{raw_type!r}", raw_preview))
                continue
            qtype = _infer_qtype(raw_answer, options_preview)      # 缺题型列 → 由答案推断
            if qtype is None:
                issues.append(BankIssue(
                    row_no, "缺少「题型」列，且无法由答案推断题型（请补题型列或填答案）", raw_preview))
                continue

        stem = _cell(row, cmap.get("stem"))
        if not stem:
            issues.append(BankIssue(row_no, "题干为空", raw_preview))
            continue
        if "[公式无缓存值]" in stem:
            issues.append(BankIssue(
                row_no, "题干单元格是公式且没有缓存值：请在 Excel 中把公式结果粘贴为值后重试",
                raw_preview))
            continue

        options = options_preview
        if qtype is not QType.JUDGE:
            letters_seq = [letter for letter, _ in texts]
            expected_seq = list(OPTION_LABELS[:len(letters_seq)])
            if letters_seq != expected_seq:
                missing = [lab for lab in expected_seq if lab not in letters_seq]
                issues.append(BankIssue(
                    row_no,
                    f"选项列不连续：缺少「选项{''.join(missing) or '?'}」列，"
                    "为避免答案与选项错位，本题未导入；请补齐选项列后重试",
                    raw_preview))
                continue
        if any(not t for t in options):
            issues.append(BankIssue(row_no, "选项不连续（中间存在空列），请整理后重新导入", raw_preview))
            continue
        if qtype is not QType.JUDGE and len(options) < 2:
            issues.append(BankIssue(
                row_no, f"{qtype.value}缺少选项（本题只识别到 {len(options)} 个选项列），"
                        "请检查列头是否写成了「选项A / 选项B …」或使用单列「选项」列", raw_preview))
            continue

        if not raw_answer and qtype is QType.JUDGE:
            issues.append(BankIssue(
                row_no, "缺少答案（若该单元格是公式，请把公式转换为值/纯文本后再导入）", raw_preview))
            continue
        if qtype is QType.JUDGE:
            # 判断题选项常写成「A正确;B错误」/「正确;错误」：先按选项文本还原答案，
            # 再把选项去掉（与段落解析、模板一致：判断题不设选项）。
            raw_answer = _resolve_judge_answer(raw_answer, options)
        try:
            letters, judge, fixed = parse_answer(qtype, raw_answer, max(len(options), 1))
        except AnswerParseError as exc:
            issues.append(BankIssue(row_no, str(exc), raw_preview))
            continue

        running += 1
        q = Question(
            qid=f"{qtype.key}-{running:04d}",
            qtype=qtype,
            stem=stem,
            options=[] if qtype is QType.JUDGE else options,
            answer_letters=letters,
            judge_answer=judge,
            fixed_order=fixed,
            analysis=_cell(row, cmap.get("analysis")),
            src_row=row_no,
        )
        problems = q.validate()
        if problems:
            issues.append(BankIssue(row_no, "；".join(problems), raw_preview))
            continue
        questions.append(q)

    return questions, issues, header, dict(cmap.items())


# ================================================================ 段落文本解析
# 用于 Word 段落型题库与 PDF 文本型题库：逐行状态机识别章节/题干/选项/答案/解析。

_FW_LETTERS = str.maketrans("ＡＢＣＤＥＦＧＨａｂｃｄｅｆｇｈ", "ABCDEFGHabcdefgh")
_FW_DIGITS = str.maketrans("０１２３４５６７８９", "0123456789")

_LETTER_CLS = "A-Ha-hＡ-Ｈａ-ｈ"

_SECTION_RE = re.compile(
    r"^\s*(?:第\s*[一二三四五六七八九十0-9０-９]+\s*部分\s*[:：]?|[一二三四五六七八九十]+\s*[、.．])"
    r"\s*(.*?)\s*$"
)
_QNUM_LOOSE_RE = re.compile(
    r"^\s*([0-9０-９]{1,6})\s*([、．.．)）])\s*(\S.*)$")
_QNUM_SPACE_RE = re.compile(r"^\s*([0-9０-９]{1,6})\s+(\S.*)$")
_QNUM_CN_RE = re.compile(r"^\s*第\s*([0-9０-９]{1,6})\s*题\s*[、.．:：]?\s*(.*)$")
_STEM_PREFIX_RE = re.compile(r"^\s*题\s*干\s*[:：]\s*(.*)$")
_ANS_KEYWORDS = r"(?:参考答案|正确答案|标准答案|答案)"
_ANA_KEYWORDS = (r"(?:答案解析|答案要点|答案说明|试题解析|题目解析|解析|分析|说明|依据|"
                 r"知识点|出处|来源|文件依据)")
_ANS_RE = re.compile(
    rf"^\s*(?:[【\[（(]\s*{_ANS_KEYWORDS}\s*[】\]）)]"
    rf"|{_ANS_KEYWORDS}(?=\s*[:：]|\s|[A-Ha-h1-8√✓×✗对错正确错误]))"
    r"\s*[:：]?\s*(\S.*)$"
)
# 解析行必须先于答案行匹配：否则「答案解析：…」会被 _ANS_RE 当成答案
# （审计 R2：纯中文解析会让整题被拒，含字母解析会静默改答案/改题型）。
_ANA_RE = re.compile(
    rf"^\s*(?:[【\[（(]\s*{_ANA_KEYWORDS}\s*[】\]）)]"
    rf"|(?:答案解析|答案要点|答案说明|试题解析|题目解析|解析|分析)(?=\s*[:：]|\s)"
    rf"|(?:说明|依据|知识点|出处|来源|文件依据)(?=\s*[:：]))"
    r"\s*[:：]?\s*(.*)$"
)
_OPT_MARK_RE = re.compile(rf"([{_LETTER_CLS}])\s*[.．、,，:：]\s*")
_OPT_START_RE = re.compile(rf"^\s*[{_LETTER_CLS}]\s*[.．、,，:：]")
_LETTER_CHARS = "ABCDEFGHabcdefghＡＢＣＤＥＦＧＨａｂｃｄｅｆｇｈ"
_CJK_RANGES = ((0x4E00, 0x9FFF), (0x3400, 0x4DBF))
_PUNCT_BEFORE_OPTION = "，。；：、）)】］>"
_BRACKET_RE = re.compile(r"[（(]\s*([^（）()]{0,16}?)\s*[）)]")
_ANSWER_TOKEN_RE = re.compile(
    rf"^(?:[{_LETTER_CLS}]{{1,8}}(?:\s*[、,，/]\s*[{_LETTER_CLS}]{{1,8}})*"
    r"|[√✓×✗对错]|正确|错误|是|否|T|F|True|False)$",
    re.IGNORECASE,
)
_GLUE_RE = re.compile(r"(?<=\S)\s*([0-9０-９]{1,6})\s*(?:[、．)）]\s*|\.\s*(?![0-9０-９]))(\S.*)$")
_PAGE_NO_RE = re.compile(
    r"^\s*(?:第\s*[0-9]{1,4}\s*页(?:\s*[，,/]?\s*共\s*[0-9]{1,4}\s*页)?"
    r"|[-—–]?\s*[0-9]{1,4}\s*[-—–]?"
    r"|[0-9]{1,4}\s*/\s*[0-9]{1,4})\s*$"
)

_SECTION_TYPES: list[tuple[tuple[str, ...], QType]] = [
    (("单项选择", "单选"), QType.SINGLE),
    (("多项选择", "多选", "不定项"), QType.MULTIPLE),
    (("判断题", "判断", "是非", "对错"), QType.JUDGE),
]
_SECTION_SKIP = ("案例", "主观", "简答", "问答", "论述", "计算", "填空", "绘图", "实操", "名词解释")


@dataclass
class TextLine:
    """段落文本中的一行（Word 段落 / 软换行 / PDF 文本行）。"""

    text: str
    is_heading: bool = False
    source: int = 0
    auto_numbered: bool = False      # Word 自动编号（编号不在 paragraph.text 里）


def _parse_question_number(text: str, expected: "int | None"
                           ) -> "tuple[int | None, str] | None":
    """识别"新题起始"行，返回 ``(题号或 None, 题干)``；不是新题返回 None。

    ``expected`` 是「上一题号 + 1」（未知为 None）。以下三种真实写法都支持：
    * ``12、题干`` / ``12. 题干`` / ``12．题干``（常规）
    * ``58.2022 版…`` / ``90.2026年…``（题号后紧跟数字，不能被当成小数）
    * ``24 门窗工程…``（题号后只有空格）
    后两种与正文里的 ``1.5 米`` 形态相似，因此只在题号等于 ``expected`` 时采信。
    """
    m = _QNUM_CN_RE.match(text)
    if m:
        stem = (m.group(2) or "").strip()
        if stem:
            return int(m.group(1)), stem
    m = _QNUM_LOOSE_RE.match(text)
    if m:
        num = int(m.group(1))
        rest = (m.group(3) or "").strip()
        if not rest:
            return None
        if rest[0].isdigit() and num != expected:
            # 「1.5 米」这类小数/正文，不是题号
            return None
        return num, rest
    m = _QNUM_SPACE_RE.match(text)
    if m:
        num = int(m.group(1))
        rest = (m.group(2) or "").strip()
        if (num == expected and rest and not rest[0].isdigit()
                and any(_is_cjk(ch) for ch in rest[:8])):
            return num, rest
    return None


def _inside_parens(text: str, pos: int) -> bool:
    """pos 是否落在未闭合的中/英文括号内。"""
    depth = 0
    for ch in text[:pos]:
        if ch in "（(":
            depth += 1
        elif ch in "）)":
            depth = max(0, depth - 1)
    return depth > 0


def _find_glued_question(text: str, expected: "int | None"):
    """在选项行尾部找"粘在上一题末尾的下一题"，返回 match 或 None。

    要求：匹配不在括号内、匹配前已有选项标记、题号等于 expected（已知时）、
    tail 不以数字开头（排除「约 5、7 米」）。
    """
    for gm in _GLUE_RE.finditer(text):
        if gm.start() == 0 or _inside_parens(text, gm.start()):
            continue
        tail = gm.group(2)
        if len(tail) < 2 or tail[0].isdigit():
            continue
        number = int(gm.group(1))
        if expected is not None and number != expected:
            # 题号不连续（真实题库里偶有跳号）：只有 tail 明显像题干时才切，
            # 避免把「约 8、9 米」「共 3、4 项」误判成下一题。
            if not (len(tail) >= 6 and re.search(r"[：:？（(。]", tail)):
                continue
        if not _iter_option_marks(text[: gm.start()]):
            continue
        return gm
    return None


def _letter(ch: str) -> str:
    return ch.translate(_FW_LETTERS).upper()


def _digits(text: str) -> str:
    return text.translate(_FW_DIGITS)


def _classify_section(title: str) -> "QType | str | None":
    """章节标题 → 题型；``"skip"`` 表示该章节不收录（主观题等）。"""
    t = title.strip()
    for key in _SECTION_SKIP:
        if key in t:
            return "skip"
    for keys, qtype in _SECTION_TYPES:
        for key in keys:
            if key in t:
                return qtype
    return None


def _has_section_keyword(title: str) -> bool:
    t = title.strip()
    if any(k in t for k in _SECTION_SKIP):
        return True
    return any(k in t for keys, _ in _SECTION_TYPES for k in keys)


def _judge_from_token(token: str) -> "bool | None":
    """判断题答案 token → True/False；不是判断型答案则返回 None。"""
    t = _digits(token).strip().upper()
    if t in ("√", "✓", "对", "正确", "是", "T", "TRUE", "A", "1"):
        return True
    if t in ("×", "✗", "x", "X", "错", "错误", "否", "F", "FALSE", "B", "0"):
        return False
    return None


def _letters_from_token(token: str) -> list[str]:
    out: list[str] = []
    for ch in token:
        letter = _letter(ch) if ch in "ABCDEFGHabcdefghＡＢＣＤＥＦＧＨａｂｃｄｅｆｇｈ" else ""
        if letter and letter not in out:
            out.append(letter)
    return sorted(out)


_JUDGE_OPTION_WORDS = ("正确", "错误", "对", "错", "是", "否", "√", "×", "✓", "✗")


def _options_look_like_judge(options: dict[str, str]) -> bool:
    """选项本身是否是"正确 / 错误"两类（据此判断字母答案是不是判断题答案）。"""
    texts = [normalize_text(t) for t in options.values() if t is not None]
    if len(texts) != 2:
        return False
    return all(any(w in t for w in _JUDGE_OPTION_WORDS) for t in texts)


def _resolve_judge_answer(raw_answer: str, options: Sequence[str]) -> str:
    """判断题答案写成选项字母时，按**选项文本**还原成「正确 / 错误」。

    题库里判断题的选项常写成 ``A正确;B错误`` 或 ``正确;错误``；若顺序被写成
    ``错误;正确``，直接把字母按「A=正确」硬解会得到相反的结论，因此这里先用
    选项文本把字母翻译回文字答案（无法翻译时原样返回，交给 ``parse_answer``）。
    """
    token = normalize_text(raw_answer)
    texts = [normalize_text(t) for t in options if normalize_text(t)]
    if len(texts) != 2 or not all(any(w in t for w in _JUDGE_OPTION_WORDS) for t in texts):
        return token
    letters = _letters_from_token(token)
    if len(letters) != 1:
        return token
    idx = OPTION_LABELS.index(letters[0])
    return texts[idx] if idx < len(texts) else token


def _extract_embedded_answer(stem: str) -> tuple["str | None", str]:
    """从题干括号里提取答案，并把答案抹成作答空位。

    只处理"像答案"的括号内容（``（C）``/``（ABCD）``/``（  BCE ）``/``（√）``），
    像 ``（含）`` 这种普通括号原样保留；空括号 ``（  ）`` 也不动。
    返回 ``(答案token 或 None, 处理后的题干)``。
    """
    found: "str | None" = None

    def _replace(m: re.Match) -> str:
        nonlocal found
        token = m.group(1).strip()
        if not token or not _ANSWER_TOKEN_RE.match(token):
            return m.group(0)
        if found is None:
            found = token
        return "（　　）"

    cleaned = _BRACKET_RE.sub(_replace, stem)
    return found, cleaned


def _is_cjk(ch: str) -> bool:
    code = ord(ch)
    return any(lo <= code <= hi for lo, hi in _CJK_RANGES)


def _is_latin(ch: str) -> bool:
    return ("A" <= ch.upper() <= "Z") or ("0" <= ch <= "9")


def _iter_option_marks(text: str) -> list[tuple[int, str, int]]:
    """扫描一行中所有"像选项标记"的位置。

    返回 ``[(标记起始下标, 选项字母, 选项文本起始下标)]``。

    识别三种常见写法：
    * ``A.内容`` / ``A．内容`` / ``A、内容``（带分隔符）
    * ``C  2.5m``（字母 + 空格，无分隔符）
    * ``A明确内容``（字母后直接接中文的写法）

    同时排除 ``ABC`` 这类正文里的拉丁字母串。
    """
    out: list[tuple[int, str, int]] = []
    n = len(text)
    for i, ch in enumerate(text):
        if ch not in _LETTER_CHARS:
            continue
        letter = _letter(ch)
        prev = text[i - 1] if i else ""
        if i and not (prev.isspace() or prev in _PUNCT_BEFORE_OPTION or _is_cjk(prev)):
            continue
        j = i + 1
        while j < n and text[j] in " \u3000\t":
            j += 1
        if j < n and text[j] in ".．、,，:：":
            j += 1
            while j < n and text[j] in " \u3000\t":
                j += 1
            if j < n:
                out.append((i, letter, j))
            continue
        if j < n and text[j] in ")）":
            out.append((i, letter, j + 1))
            continue
        if j < n and not ("A" <= text[j].upper() <= "Z"):
            out.append((i, letter, j))
    return out


def _is_option_line(text: str) -> bool:
    if len(text) > 400:
        return False
    marks = _iter_option_marks(text)
    if not marks:
        return False
    return marks[0][0] == 0 or len(marks) >= 2


def _split_options(text: str, expected: "int | None" = None,
                   allow_glue: bool = True
                   ) -> tuple[list[tuple[str, str]], "str | None"]:
    """把一行切成 ``[(字母, 选项文本)]``，并返回尾部粘连的下一题内容。

    粘连题必须**先于选项切分**识别：粘连文本里常带 ``（ C ）`` 形式的答案，
    否则括号里的字母会被 :func:`_iter_option_marks` 当成额外选项，导致
    下一题被吞、上一题选项被截断（审计 P1-1）。
    """
    leftover: "str | None" = None
    if allow_glue:
        glue = _find_glued_question(text, expected)
        if glue is not None:
            leftover = f"{_digits(glue.group(1))}、{glue.group(2)}"
            text = text[: glue.start()].rstrip()

    marks = _iter_option_marks(text)
    if not marks:
        return [], leftover

    segs: list[tuple[str, str]] = []
    for i, (_start, letter, content_start) in enumerate(marks):
        end = marks[i + 1][0] if i + 1 < len(marks) else len(text)
        segs.append((letter, text[content_start:end].strip()))
    return [(lab, txt) for lab, txt in segs if txt], leftover


# 单列「选项」单元格里的显式分隔符：`;`/`；`/`|`/换行，
# 以及题库模板导出常见的 `$;$`（先匹配，避免 `$` 参与后续切分）。
_OPT_CELL_SEP_RE = re.compile(r"\$;\$|[;；|｜\n\r]+")
# 强字母标记：字母后有分隔符或空格（`A、F15` / `A.甲` / `A 2/3` / `A)甲`），基本不会误判。
_OPT_LABEL_STRONG_RE = re.compile(rf"^([{_LETTER_CLS}])(?:[.．、,，:：)）]\s*|\s+)")
# 弱字母标记：字母后直接接非拉丁字符（`A绑扎连接` / `A正确`）。
# 这类写法有歧义（`A票由…` 里的 A 属于正文），只在整列字母自洽时才采信。
_OPT_LABEL_WEAK_RE = re.compile(rf"^([{_LETTER_CLS}])(?=[^\sA-Za-z0-9])")


def _common_prefix(texts: Sequence[str]) -> str:
    """一组文本的公共前缀（任一为空则返回空串）。"""
    if not texts:
        return ""
    first = texts[0]
    for other in texts[1:]:
        n = 0
        limit = min(len(first), len(other))
        while n < limit and first[n] == other[n]:
            n += 1
        first = first[:n]
        if not first:
            return ""
    return first


def _split_option_cell(cell: str) -> list[tuple[str, str]]:
    """把「选项写在一列」的单元格切成 ``[(字母, 选项文本)]``。

    比 :func:`_split_options` 多支持两类真实题库写法：

    * **显式分隔符**：``A、F15;B、F25``（``;``/``；``/``|``/换行，以及模板导出的 ``$;$``）
    * **完全不带字母**：``正确;错误`` / ``甲;乙;丙;丁`` —— 按出现顺序补 A、B、C…

    归位规则（按优先级）：

    1. 每个片段都带字母、且字母严格递增 ⇒ 整列按字母归位（``A绑扎连接;B焊接``）；
    2. 有 ≥2 个「字母 + 分隔符/空格」的强标记 ⇒ 强标记按字母归位，其余片段按顺序补字母
       （源文件漏写字母如 ``A.雨后;持续淋水1h;C.…`` 时也能还原出 B）；
    3. 其余情况 ⇒ 整列按顺序编号，片段正文原样保留（``A票由…`` 里的 A 属于正文）。

    没有显式分隔符时回退到 :func:`_split_options`，继续支持 ``A.甲 B.乙`` 这类写法。
    """
    text = normalize_display(cell)
    if not text:
        return []
    parts = [p for p in (s.strip() for s in _OPT_CELL_SEP_RE.split(text)) if p]
    if len(parts) < 2:
        return _split_options(text, allow_glue=False)[0]

    parsed: list[tuple[str, str, str, str]] = []      # (强字母, 弱字母, 强字母后的正文, 原片段)
    for part in parts:
        m = _OPT_LABEL_STRONG_RE.match(part)
        if m:
            parsed.append((_letter(m.group(1)), "", part[m.end():].strip(), part))
            continue
        m = _OPT_LABEL_WEAK_RE.match(part)
        if m:
            parsed.append(("", _letter(m.group(1)), part[m.end():].strip(), part))
            continue
        parsed.append(("", "", "", part))

    # 弱字母歧义：「A类火灾;B类火灾」/「A票由…;B票由…」里的 A/B 是正文而不是选项标签。
    # 若所有弱标记片段的剩余正文有公共前缀（含"完全相同"），则不把字母当标签。
    weak_rests = [rest for _s, weak, rest, _f in parsed if weak and rest]
    if len(weak_rests) >= 2 and _common_prefix(weak_rests):
        parsed = [(strong, "", rest, full) for strong, _weak, rest, full in parsed]

    labelled = [strong or weak for strong, weak, _rest, _full in parsed]
    # ① 全部片段都带字母、且字母严格递增 ⇒ 按字母归位
    if all(labelled) and all(labelled[i] < labelled[i + 1] for i in range(len(labelled) - 1)):
        return [(strong or weak, rest) for strong, weak, rest, _full in parsed if rest]

    # ② 至少 2 个强标记 ⇒ 强标记按字母归位，其余片段补下一个未用字母
    if sum(1 for strong, _w, _r, _f in parsed if strong) >= 2:
        used = {strong for strong, _w, _r, _f in parsed if strong}
        cursor = 0
        out: list[tuple[str, str]] = []
        for strong, _weak, rest, full in parsed:
            if strong:
                out.append((strong, rest))
                continue
            while cursor < len(OPTION_LABELS) and OPTION_LABELS[cursor] in used:
                cursor += 1
            label = OPTION_LABELS[cursor] if cursor < len(OPTION_LABELS) else ""
            used.add(label)
            out.append((label, full))
        return [(label, txt) for label, txt in out if label and txt]

    # ③ 整列没有可信字母 ⇒ 按顺序编号，正文原样保留
    return [(OPTION_LABELS[i] if i < len(OPTION_LABELS) else "", full)
            for i, (_strong, _weak, _rest, full) in enumerate(parsed) if full]


def _decide_qtype(section: "QType | str | None", answer_token: str,
                  options: dict[str, str]) -> "QType | None":
    """题型判定：章节优先，答案形态兜底；单选章节出现多字母答案时容错为多选。

    注意：字母答案 ``A``/``B`` 既可能是单选选项，也可能是判断题的"正确/错误"，
    因此**只有在没有选项、或选项本身就是"正确/错误"两类**时才按判断题处理，
    否则 ``（B）`` 这类单选答案会被误判成判断题并丢掉全部选项。
    """
    letters = _letters_from_token(answer_token)
    judge = _judge_from_token(answer_token) if answer_token else None
    if section is QType.JUDGE:
        return QType.JUDGE
    judge_like = (not options) or _options_look_like_judge(options)
    if judge is not None and judge_like:
        return QType.JUDGE
    if section is QType.MULTIPLE:
        return QType.MULTIPLE
    if section is QType.SINGLE:
        return QType.MULTIPLE if len(letters) >= 2 else QType.SINGLE
    if len(letters) >= 2:
        return QType.MULTIPLE
    if len(letters) == 1:
        return QType.SINGLE
    if options:
        return QType.SINGLE
    return None


def lines_to_questions(lines: Sequence[TextLine], *, start_line: int = 1
                       ) -> tuple[list[Question], list[BankIssue], dict]:
    """解析段落型题库（Word 段落 / PDF 文本）。

    返回 ``(questions, issues, meta)``；``meta`` 含 ``warnings`` 与计数信息。
    """
    buf: "deque[TextLine]" = deque(lines)
    questions: list[Question] = []
    issues: list[BankIssue] = []
    warnings: list[str] = []
    counter = 0
    section: "QType | str | None" = None
    section_name = ""
    section_serial = 0
    skipped_sections = 0
    expected: "int | None" = None
    qnums_by_section: dict[str, list[int]] = {}
    cur: "dict | None" = None
    row_no = start_line - 1

    def flush() -> None:
        nonlocal cur, counter
        if cur is None:
            return
        data, cur = cur, None
        if data.get("qnum") is not None:
            key = f"{data.get('section_serial', 0)}:{data.get('section_name', '')}"
            qnums_by_section.setdefault(key, []).append(int(data["qnum"]))
        stem = re.sub(r"\s+", " ", "".join(data["stem"])).strip()
        if not stem:
            return
        if "[公式无缓存值]" in stem:
            issues.append(BankIssue(data["row"], "题干单元格是公式且没有缓存值，"
                                    "请在 Excel 中把公式结果粘贴为值后重试", stem[:120]))
            return
        answer_token = data["answer"].strip()
        if not answer_token:
            token, cleaned = _extract_embedded_answer(stem)
            if token:
                answer_token = token
                stem = re.sub(r"\s+", " ", cleaned).strip()
        sec = data["section"]
        if sec == "skip":
            return
        options = data["options"]
        qtype = _decide_qtype(sec, answer_token, options)
        if qtype is None:
            issues.append(BankIssue(data["row"], "未找到答案（题干括号为空且无答案行）",
                                    stem[:120]))
            return
        text_letters = _letters_from_token(answer_token)
        need_options = qtype is not QType.JUDGE
        if need_options and len(options) < 2:
            issues.append(BankIssue(data["row"], f"选项少于 2 个（实际 {len(options)}）", stem[:120]))
            return
        if not answer_token:
            issues.append(BankIssue(
                data["row"], "缺少答案（若该单元格是公式，请把公式转换为值/纯文本后再导入）",
                stem[:120]))
            return
        # 单选章节里出现了多字母答案 → 容错为多选并提示
        if sec is QType.SINGLE and len(text_letters) >= 2:
            warnings.append(f"第 {data['row']} 行：章节标注为单选题，但答案有 {len(text_letters)} 个字母，"
                            f"已按多选题收录。")
        try:
            letters, judge_value, fixed = parse_answer(
                qtype, answer_token, max(len(options), 1))
        except AnswerParseError as exc:
            issues.append(BankIssue(data["row"], str(exc), stem[:120]))
            return
        counter += 1
        ordered = sorted(options.items())
        q = Question(
            qid=f"{qtype.key}-{counter:04d}",
            qtype=qtype,
            stem=stem,
            options=[text for _, text in ordered] if need_options else [],
            answer_letters=letters,
            judge_answer=judge_value,
            fixed_order=fixed,
            analysis=" ".join(a for a in data["analysis"] if a).strip(),
            src_row=data["row"],
        )
        problems = q.validate()
        if problems:
            issues.append(BankIssue(data["row"], "；".join(problems), stem[:120]))
            counter -= 1
            return
        questions.append(q)

    def start_question(stem: str, qnum: "int | None", auto: bool = False) -> None:
        nonlocal cur, expected
        flush()
        cur = {"stem": [stem], "options": {}, "answer": "", "analysis": [],
               "row": row_no, "section": section, "section_name": section_name,
               "section_serial": section_serial, "qnum": qnum}
        expected = (qnum + 1) if qnum is not None else None
        if auto is False and qnum is None:
            expected = None

    while buf:
        ln = buf.popleft()
        text = (ln.text or "").strip()
        if not text:
            continue
        row_no = ln.source or (row_no + 1)

        # 1) 章节标题（题型切换）：支持「一、单项选择题」与 Word 标题样式（简答题等）
        m = _SECTION_RE.match(text)
        title = m.group(1) if m else (text if ln.is_heading else "")
        if title and _has_section_keyword(title):
            flush()
            section_name = str(title).strip()
            section = _classify_section(section_name)
            section_serial += 1
            expected = None
            if section == "skip":
                skipped_sections += 1
            continue

        # 2) 解析/来源行（必须先于答案行：避免「答案解析：」误当答案，审计 R2）
        m = _ANA_RE.match(text)
        if m:
            if cur is not None:
                cur["analysis"].append((m.group(1) or "").strip())
            continue

        # 3) 答案行
        m = _ANS_RE.match(text)
        if m:
            if cur is not None:
                cur["answer"] = (m.group(1) or "").strip()
            continue

        # 4) 新题起始（常规题号 / 题号后接数字 / 题号后空格 / 第 N 题）
        parsed = _parse_question_number(text, expected)
        if parsed is not None:
            qnum, stem = parsed
            start_question(stem, qnum)
            continue
        m = _STEM_PREFIX_RE.match(text)
        if m:
            start_question((m.group(1) or "").strip(), None)
            continue

        # 5) 选项行（粘连的下一题先切出来，见 _split_options）
        if _is_option_line(text) and cur is not None:
            segs, leftover = _split_options(text, expected=expected)
            for letter, opt in segs:
                if letter not in cur["options"]:
                    cur["options"][letter] = opt
            if leftover:
                buf.appendleft(TextLine(leftover, source=row_no))
            continue

        # 5b) Word 自动编号段落：编号不在 paragraph.text 里，但每个编号段落就是一道新题
        if (ln.auto_numbered and section not in (None, "skip") and cur is not None
                and len(text) >= 4 and not _is_option_line(text)):
            start_question(text, None, auto=True)
            continue
        if (ln.auto_numbered and section not in (None, "skip") and cur is None
                and len(text) >= 4 and not _is_option_line(text)):
            start_question(text, None, auto=True)
            continue

        # 6) 续行：多行题干 / 折行的选项 / 说明
        if cur is None:
            # 首题没有编号：出现空括号/问号且有选项或答案时，当作题干起点（审计 R15）
            if (section not in (None, "skip")
                    and (re.search(r"[（(]\s*[）)]", text) or text.endswith(("？", "?")))
                    and len(text) >= 5):
                start_question(text, None)
            continue
        if not cur["options"]:
            cur["stem"].append(text)
        elif not cur["answer"] and not cur["analysis"] and cur["options"]:
            last = sorted(cur["options"])[-1]           # 选项折行：接到最后一个选项
            cur["options"][last] += text
        else:
            cur["analysis"].append(text)

    flush()

    # 题号连续性自检：接近连续编号的章节若缺号，提示"可能被并入相邻题目"，避免静默丢题
    for key, nums in qnums_by_section.items():
        uniq = sorted(set(nums))
        if len(uniq) < 2:
            continue
        lo, hi = uniq[0], uniq[-1]
        if hi - lo + 1 > len(uniq) + 20:
            continue                                  # 全局跨章节编号，不做连续性判断
        missing = [n for n in range(lo, hi + 1) if n not in set(uniq)]
        if missing:
            name = key.split(":", 1)[-1] or "(开头)"
            sample = "、".join(str(n) for n in missing[:10])
            if len(missing) > 10:
                sample += "…"
            warnings.append(f"章节「{name}」题号不连续：缺少 {sample}（可能被并入相邻题目）")

    meta = {"warnings": warnings, "total_lines": len(lines),
            "skipped_sections": skipped_sections}
    return questions, issues, meta


# ---------------------------------------------------------------- Word / PDF
def _read_docx(path: Path) -> tuple[list[list[list[str]]], list[TextLine]]:
    """读取 Word：返回 (若干表格, 段落行)。"""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - 仅在缺依赖时发生
        raise ValueError("缺少 python-docx，无法读取 Word 题库，请重新安装或改用 Excel。") from exc

    try:
        doc = Document(str(path))
    except Exception as exc:
        raise ValueError(f"无法打开 Word 文件（可能已损坏或为 .doc 旧格式）：{exc}") from exc

    tables: list[list[list[str]]] = []
    for table in doc.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            # 合并单元格在 row.cells 里会重复出现：保留网格列位置（文本占用首列，
            # 其余列补空），否则整行会左移、答案列错位（审计 R21）。
            cells: list[str] = []
            prev_key: "int | None" = None
            span = 0
            for cell in row.cells:
                key = id(cell._tc)
                if key == prev_key:
                    span += 1
                    continue
                if prev_key is not None and span > 1:
                    cells.extend([""] * (span - 1))
                cells.append(normalize_display(cell.text))
                prev_key = key
                span = 1
            if prev_key is not None and span > 1:
                cells.extend([""] * (span - 1))
            rows.append(cells)
        if rows:
            tables.append(rows)

    lines: list[TextLine] = []
    for idx, para in enumerate(doc.paragraphs, start=1):
        raw = para.text or ""
        if not raw.strip():
            continue
        style_name = ""
        try:
            style_name = str(para.style.name or "")
        except Exception:
            style_name = ""
        is_heading = ("heading" in style_name.lower()) or ("标题" in style_name)
        auto_numbered = False
        try:
            ppr = para._p.pPr
            auto_numbered = bool(ppr is not None and ppr.numPr is not None)
        except Exception:
            auto_numbered = False
        for part in raw.split("\n"):
            part = part.strip()
            if part:
                lines.append(TextLine(part, is_heading=is_heading, source=idx,
                                      auto_numbered=auto_numbered))
    return tables, lines


def _read_pdf(path: Path) -> list[TextLine]:
    """读取 PDF 文本行；无文字层时给出可读提示。"""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ValueError("缺少 pypdf，无法读取 PDF 题库，请重新安装或改用 Word/Excel。") from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise ValueError(f"无法打开 PDF 文件（可能已损坏或格式不受支持）：{exc}") from exc

    if getattr(reader, "is_encrypted", False):
        try:
            if reader.decrypt("") == 0:
                raise ValueError("PDF 已加密")
        except Exception as exc:
            raise ValueError("PDF 已加密，无法读取；请先解除密码保护或改用 Word/Excel 题库。") from exc

    lines: list[TextLine] = []
    try:
        pages = list(reader.pages)
    except Exception as exc:
        raise ValueError(f"无法解析 PDF 页面：{exc}") from exc

    page_lines: list[list[str]] = []
    for page in pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        current: list[str] = []
        for raw in text.splitlines():
            raw = normalize_display(raw)
            if raw:
                current.append(raw)
        page_lines.append(current)

    # 页眉/页脚：在多页重复出现的同一行（且不太长）视为页眉页脚，予以剔除；
    # 纯页码行（"3" / "第 3 页" / "3/10"）直接丢弃（审计 R10）。
    counter: "Counter[str]" = Counter()
    for current in page_lines:
        for raw in set(current):
            if len(raw) <= 80:
                counter[normalize_text(raw)] += 1
    repeat_threshold = max(2, int(len(page_lines) * 0.6))
    repeated = {t for t, n in counter.items() if n >= repeat_threshold} if len(page_lines) >= 2 else set()
    line_no = 0
    for _pno, current in enumerate(page_lines, start=1):
        for raw in current:
            if _PAGE_NO_RE.match(raw):
                continue
            if normalize_text(raw) in repeated:
                continue
            line_no += 1
            # 用"页内累计行号"而不是页码作 source：同页多题的行号不再相同（审计 R25）
            lines.append(TextLine(raw, source=line_no))

    if not lines:
        raise ValueError("未能从 PDF 中提取到任何文字：该文件很可能是扫描件（图片），"
                         "请改用文本版 PDF，或提供 Word/Excel 题库。")
    return lines


def _pdf_lines_to_rows(lines: Sequence[TextLine]) -> list[list[str]]:
    """PDF 表格化兜底：按连续多空格分列。"""
    rows: list[list[str]] = []
    for ln in lines:
        cells = [c.strip() for c in re.split(r"\s{2,}|\t", ln.text) if c.strip()]
        if len(cells) >= 2:
            rows.append(cells)
    return rows


_COMBINED_OPTION_HEADERS = {"选项", "选项内容", "备选项", "答案选项", "可选项", "选项列表",
                            "试题选项", "题目选项", "试题选项内容",
                            "options", "option", "选项（一行）"}


def _expand_combined_options(rows: list[list[str]]) -> list[list[str]]:
    """把"所有选项挤在一列"的表格展开成 选项A…选项H 多列。

    表头不一定在第 1 行（上方可能有标题/空行），因此先用与列头识别相同的规则定位表头行。
    """
    if not rows:
        return rows

    stem_tokens = {_header_token(n) for n in _HEADER_SYNONYMS["stem"]}
    header_idx = -1
    combo = None
    for i, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        tokens = [_header_token(c) for c in row]
        if not any(t in _COMBINED_OPTION_HEADERS for t in tokens):
            continue
        if not any(t in stem_tokens for t in tokens if t):
            continue
        header_idx = i
        combo = next(j for j, t in enumerate(tokens) if t in _COMBINED_OPTION_HEADERS)
        break
    if header_idx < 0 or combo is None:
        return rows

    header = rows[header_idx]
    if any(_option_letter_of_header(h) for h in header if h):
        return rows                                        # 已有 A–H 独立列
    # 两行表头：「选项」组标题 + 下一行 A B C D 子表头 → 交给 rows_to_questions
    # 合并子表头，绝不能把数据行当"单列选项"切分（审计 R6）。
    if header_idx + 1 < len(rows) and _option_subheader_letters(rows[header_idx + 1]) is not None:
        return rows

    letters = list(OPTION_LABELS[:8])
    out = [list(r) for r in rows[:header_idx]]              # 标题/空行原样保留
    out.append(list(header[:combo]) + [f"选项{x}" for x in letters] + list(header[combo + 1:]))
    for row in rows[header_idx + 1:]:
        cell = row[combo] if combo < len(row) else ""
        opts = dict(_split_option_cell(cell))
        out.append(list(row[:combo]) + [opts.get(x, "") for x in letters] + list(row[combo + 1:]))
    return out


def _best_of_tables(tables: Sequence[list[list[str]]]) -> tuple[list[Question], list[BankIssue], list[str], dict]:
    """多张 Word 表格时，取能解析出最多题目的那一个。

    解析出 0 题的表格也保留其坏行原因（用于"整篇 0 题"时的可读报错）。
    """
    best: "tuple[list[Question], list[BankIssue], list[str], dict] | None" = None
    for rows in tables:
        for candidate in (rows, _expand_combined_options(rows)):
            try:
                qs, issues, header, cmap = rows_to_questions(candidate)
            except Exception:
                continue
            if best is None or (len(qs), len(issues)) > (len(best[0]), len(best[1])):
                best = (qs, issues, header, cmap)
    return best if best is not None else ([], [], [], {})


def load_bank(path: str | Path) -> BankLoadResult:
    """从文件加载题库。"""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"题库文件不存在：{p}")
    suffix = p.suffix.lower()
    if suffix == ".xls":
        raise ValueError("暂不支持 .xls（Excel 97-2003 旧格式）。请用 Excel/WPS 打开后"
                         "「另存为 → Excel 工作簿(*.xlsx)」再导入。")
    if suffix == ".doc":
        raise ValueError("暂不支持 .doc（Word 97-2003 旧格式）。请用 Word 另存为 "
                         "「Word 文档(*.docx)」再导入。")
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"暂不支持的题库格式：{suffix}（支持 {'、'.join(SUPPORTED_SUFFIXES)}）")

    sheet = ""
    questions: list[Question] = []
    issues: list[BankIssue] = []
    header: list[str] = []
    cmap: dict = {}
    extra_warnings: list[str] = []

    if suffix in (".xlsx", ".xlsm"):
        sheet, rows = _read_xlsx(p)
        questions, issues, header, cmap = rows_to_questions(rows)
    elif suffix == ".json":
        rows = _read_json(p)
        questions, issues, header, cmap = rows_to_questions(rows)
    elif suffix in TEXT_SUFFIXES:
        rows = _read_delimited(p)
        questions, issues, header, cmap = rows_to_questions(rows)
    elif suffix in WORD_SUFFIXES:
        tables, lines = _read_docx(p)
        sheet = "段落文本" if lines else ""
        # 段落型与表格型都试一遍，取题目更多的那个
        q_text, i_text, meta_text = lines_to_questions(lines)
        q_tab, i_tab, h_tab, c_tab = _best_of_tables(tables)
        if len(q_tab) > len(q_text):
            questions, issues, header, cmap = q_tab, i_tab, h_tab, c_tab
            sheet = "表格"
        else:
            questions, issues = q_text, i_text
            extra_warnings = meta_text.get("warnings", [])
        if not questions:
            detail = ""
            if i_tab:
                detail = " 首个原因：" + "；".join(str(x) for x in i_tab[:3])
            raise ValueError(
                "未能从 Word 文档中识别出题目。请确认文档包含“题型/题干/选项/答案”内容；"
                "若题目是图片，需要先转成文字。" + detail)
    elif suffix in PDF_SUFFIXES:
        lines = _read_pdf(p)
        sheet = "PDF 文本"
        q_text, i_text, meta_text = lines_to_questions(lines)
        extra_warnings = meta_text.get("warnings", [])
        if not q_text:                                   # 表格化兜底
            try:
                q_tab, i_tab, h_tab, c_tab = rows_to_questions(_pdf_lines_to_rows(lines))
            except Exception:
                q_tab, i_tab, h_tab, c_tab = [], [], [], {}
            questions, issues = q_tab, i_tab
            header, cmap, sheet = h_tab or [], c_tab or {}, "PDF 表格"
        else:
            questions, issues = q_text, i_text
        if not questions:
            raise ValueError(
                "未能从 PDF 中识别出题目：可能是扫描件（图片）或版式过于特殊。"
                "请改用文本版 PDF、Word 或 Excel 题库。")
    else:                                                # pragma: no cover - 前面已拦截
        raise ValueError(f"暂不支持的题库格式：{suffix}")

    # 结果必须始终是 list[Question]（历史缺陷：PDF 兜底曾把 4 元组当题目列表）
    if not isinstance(questions, list) or any(not isinstance(q, Question) for q in questions):
        raise ValueError("题库解析结果异常（内部错误），请把该文件反馈给开发者。")
    if not questions and not issues:
        issues = [BankIssue(0, "没有解析到任何题目：请检查列头（题型/题干/选项/答案）与内容格式")]

    result = BankLoadResult(questions=questions, issues=issues, source=str(p),
                            sheet=sheet, header=header, column_map=cmap)
    if extra_warnings:
        result.column_map = dict(cmap or {})
        result.column_map.setdefault("warnings", extra_warnings)
    return result


# ---------------------------------------------------------------- 统计/导出
def bank_stats(questions: Iterable[Question]) -> dict:
    """统计题库结构。"""
    qs = list(questions)
    counts = {QType.SINGLE: 0, QType.MULTIPLE: 0, QType.JUDGE: 0}
    fixed = 0
    for q in qs:
        counts[q.qtype] += 1
        if q.fixed_order:
            fixed += 1
    return {
        "total": len(qs),
        "single": counts[QType.SINGLE],
        "multiple": counts[QType.MULTIPLE],
        "judge": counts[QType.JUDGE],
        "fixed": fixed,
    }


def questions_to_rows(questions: Iterable[Question]) -> list[list[str]]:
    """把题目转换为标准表格行（含列头），用于导出回 xlsx。"""
    header = ["题型", "题目标题", "选项A", "选项B", "选项C", "选项D", "选项E", "选项F",
              "选项G", "选项H", "解析", "答案"]
    rows = [header]
    for q in questions:
        opts = list(q.options) + [""] * (8 - len(q.options))
        answer = q.answer_display + ("（定）" if q.fixed_order else "")
        rows.append([q.qtype.value, q.stem, *opts[:8], q.analysis, answer])
    return rows


def export_questions_to_xlsx(questions: Iterable[Question], path: str | Path) -> Path:
    """把解析后的题库导出为标准 xlsx（可用于规范化用户的题库）。"""
    from openpyxl import Workbook

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "题库"
    for row in questions_to_rows(questions):
        ws.append(row)
    widths = [10, 60, 18, 18, 18, 18, 18, 18, 18, 18, 40, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    wb.save(out)
    return out


def write_template_xlsx(path: str | Path, *, samples: int = 6) -> Path:
    """写出带示例行的题库模板。"""
    qs = build_sample_questions()[:samples]
    return export_questions_to_xlsx(qs, path)


def build_sample_questions() -> list[Question]:
    """内置示例题库（同时用于自测与用户上手）。"""
    raw = [
        ("单选题", "未经论证压缩合同约定工期，属于哪类重大事故隐患？",
         ["安全管理", "质量管理", "环境保护", "消防管理"],
         "依据判定标准，属于安全管理类。", "A"),
        ("单选题", "起重机械的安全装置不全、失效，属于哪类隐患？",
         ["安全管理", "起重机械安装拆卸", "高处作业", "施工用电"],
         "属于起重机械安装拆卸类。", "B"),
        ("多选题", "以下哪些属于安全管理类重大事故隐患？",
         ["未经论证压缩合同约定工期", "发包给不具备安全生产条件的企业",
          "无资质承揽电力建设工程", "转包或违法分包主体工程"],
         "以上均属安全管理类。", "ABCD"),
        ("多选题", "构支架吊装时，当柱脚接近杯底时，应从柱四周向杯口放入（ ）个木楔，"
                   "同时收紧四周缆风绳，确认缆风绳全部固定并使立柱基本（ ）后，才能松大钩。",
         ["3～4", "4～5", "垂直", "平行"],
         "两空依次为 4～5 与垂直，答案顺序固定。", "BC（定）"),
        ("判断题", "未经论证压缩合同约定工期属于重大事故隐患。", [],
         "属于重大事故隐患。", "正确"),
        ("判断题", "可以在电缆沟内充装易燃易爆危险品。", [],
         "禁止在电缆沟内充装易燃易爆危险品。", "错误"),
    ]
    questions: list[Question] = []
    for i, (t, stem, opts, analysis, ans) in enumerate(raw, start=1):
        qtype = parse_qtype(t)
        letters, judge, fixed = parse_answer(qtype, ans, len(opts) or 1)
        questions.append(Question(qid=f"{qtype.key}-{i:04d}", qtype=qtype, stem=stem,
                                  options=list(opts), answer_letters=letters,
                                  judge_answer=judge, fixed_order=fixed,
                                  analysis=analysis, src_row=i + 1))
    return questions


def make_txt_sample() -> str:
    """示例 txt 题库内容（制表符分隔）。"""
    lines = ["\t".join(["题型", "题目标题", "选项A", "选项B", "选项C", "选项D", "解析", "答案"])]
    for q in build_sample_questions():
        opts = list(q.options) + ["", "", "", ""]
        lines.append("\t".join([q.qtype.value, q.stem, *opts[:4], q.analysis,
                                q.answer_display + ("（定）" if q.fixed_order else "")]))
    return "\n".join(lines) + "\n"
